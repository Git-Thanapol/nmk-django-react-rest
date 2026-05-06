import os
import zipfile
from io import BytesIO
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from django.conf import settings
from django.core.files.base import ContentFile

TEMPLATE_PATH = os.path.join(settings.BASE_DIR.parent, 'TWI50', 'TWI50_OFFICIAL_TEMPLATE.xlsx')

# Maps income_type code → Excel row in the income table
INCOME_ROW_MAP = {
    '1':  23,
    '2':  24,
    '3':  25,
    '4a': 26,
    '4b': 27,
    '6':  46,
}

_RELS_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
_CT_NS   = 'http://schemas.openxmlformats.org/package/2006/content-types'

# Files openpyxl drops that we must restore from the template
_RESTORE_PREFIXES = ('xl/drawings/', 'xl/media/', 'xl/charts/', 'xl/printerSettings/')
_SHEET_RELS       = 'xl/worksheets/_rels/sheet1.xml.rels'
_CONTENT_TYPES    = '[Content_Types].xml'


def _populate_sheet(ws, cert):
    company = cert.company
    vendor  = cert.vendor

    # Book / document numbers
    ws['Q2'] = cert.book_number or ''
    ws['Q3'] = cert.cert_number or ''

    # Payer (our company)
    ws['P5'] = company.national_id or ''
    ws['P6'] = company.tax_id or ''
    ws['C6'] = company.name or ''
    ws['C8'] = company.address or ''

    # Payee (vendor)
    ws['P11'] = vendor.national_id or ''
    ws['P12'] = vendor.tax_id or ''
    ws['C12'] = vendor.name or ''
    ws['C14'] = vendor.address or ''

    # Sequence + PND 3 checkbox
    ws['D16'] = cert.sequence_no or ''
    ws['H18'] = '/'

    # Income row (by selected income type)
    income_type = cert.income_type or '6'
    row = INCOME_ROW_MAP.get(income_type, 46)

    if income_type == '6':
        ws['E46'] = cert.income_description or ''

    ws[f'M{row}'] = cert.date_issued
    ws[f'O{row}'] = float(cert.amount_before_tax)
    ws[f'Q{row}'] = float(cert.tax_amount)

    # Provident fund / social security
    ws['Q52'] = float(cert.provident_fund_amount or 0)
    ws['J53'] = float(cert.social_security_amount or 0)
    ws['P54'] = cert.social_security_id or ''

    # Payer status: (1) หักภาษี ณ ที่จ่าย
    ws['A56'] = '/'

    # Signature block
    ws['H58'] = company.name or ''
    ws['G59'] = cert.date_issued.strftime('%d/%m/%Y') if cert.date_issued else ''


# ---------------------------------------------------------------------------
# Drawing restoration
#
# openpyxl drops xl/drawings/, xl/worksheets/_rels/sheet1.xml.rels, and
# xl/printerSettings/ when it re-serialises the workbook.  Without the _rels
# file Excel cannot find the drawing even if the drawing file is present.
# We restore all three from the original template zip.
# ---------------------------------------------------------------------------

def _merge_rels(mod_data: bytes, tmpl_data: bytes) -> bytes:
    """Merge drawing/media relationships from template into modified sheet rels."""
    ET.register_namespace('', _RELS_NS)
    try:
        mod_root  = ET.fromstring(mod_data)
        tmpl_root = ET.fromstring(tmpl_data)
    except ET.ParseError:
        return mod_data

    drawing_kw = ('../drawings/', '../media/', '../charts/', '../printerSettings/')
    existing   = {el.get('Target', '') for el in mod_root}
    for rel in tmpl_root:
        target = rel.get('Target', '')
        if any(kw in target for kw in drawing_kw) and target not in existing:
            mod_root.append(rel)

    return ET.tostring(mod_root, encoding='unicode').encode('utf-8')


def _merge_content_types(mod_data: bytes, tmpl_data: bytes) -> bytes:
    """Merge drawing/media Override entries from template into content types."""
    ET.register_namespace('', _CT_NS)
    try:
        mod_root  = ET.fromstring(mod_data)
        tmpl_root = ET.fromstring(tmpl_data)
    except ET.ParseError:
        return mod_data

    drawing_kw     = ('/drawings/', '/media/', '/charts/', '/printerSettings/')
    existing_parts = {el.get('PartName', '') for el in mod_root}
    for el in tmpl_root:
        part = el.get('PartName', '')
        if any(kw in part for kw in drawing_kw) and part not in existing_parts:
            mod_root.append(el)

    # Also copy Default entries (e.g. bin mime type) that openpyxl drops
    existing_exts = {el.get('Extension', '') for el in mod_root}
    for el in tmpl_root:
        ext = el.get('Extension', '')
        if ext and ext not in existing_exts:
            mod_root.append(el)

    return ET.tostring(mod_root, encoding='unicode').encode('utf-8')


def _restore_drawings(raw_bytes: bytes) -> bytes:
    """
    Patch the openpyxl-saved workbook to restore everything it dropped.

    Confirmed dropped by openpyxl (verified by inspection):
      - xl/drawings/drawing1.xml          → the actual shape definitions
      - xl/worksheets/_rels/sheet1.xml.rels → the link sheet → drawing
      - xl/printerSettings/printerSettings1.bin

    Without the _rels file Excel cannot locate the drawing even if the file
    exists, so we must always write it — whether or not it was in the saved zip.
    """
    with zipfile.ZipFile(TEMPLATE_PATH, 'r') as tz:
        tmpl_names    = set(tz.namelist())
        restore_files = {n: tz.read(n) for n in tmpl_names
                         if any(n.startswith(p) for p in _RESTORE_PREFIXES)}
        tmpl_rels     = tz.read(_SHEET_RELS)    if _SHEET_RELS    in tmpl_names else None
        tmpl_ct       = tz.read(_CONTENT_TYPES) if _CONTENT_TYPES in tmpl_names else None

    if not restore_files and not tmpl_rels:
        return raw_bytes

    output = BytesIO()
    with zipfile.ZipFile(BytesIO(raw_bytes), 'r') as mz, \
         zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as oz:

        mod_names = set(mz.namelist())

        for name in mz.namelist():
            if name in restore_files:
                # Always use the template version — openpyxl may have mangled it
                oz.writestr(name, restore_files[name])
            elif name == _SHEET_RELS and tmpl_rels:
                # Exists in saved zip → merge drawing rels in
                oz.writestr(name, _merge_rels(mz.read(name), tmpl_rels))
            elif name == _CONTENT_TYPES and tmpl_ct:
                oz.writestr(name, _merge_content_types(mz.read(name), tmpl_ct))
            else:
                oz.writestr(name, mz.read(name))

        # Add restore_files that openpyxl omitted entirely
        for name, data in restore_files.items():
            if name not in mod_names:
                oz.writestr(name, data)

        # CRITICAL: _rels file was completely absent from the saved zip.
        # Without it Excel cannot find the drawing at all.
        if _SHEET_RELS not in mod_names and tmpl_rels:
            oz.writestr(_SHEET_RELS, tmpl_rels)

    return output.getvalue()


def generate_wht_xlsx(cert):
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active
    # Keep the template's original sheet name ("TWI50") — renaming can break
    # references inside the drawing XML and printer settings
    _populate_sheet(ws, cert)

    raw = BytesIO()
    wb.save(raw)

    patched = _restore_drawings(raw.getvalue())

    filename = f"50Tawi_{cert.cert_number.replace('/', '-')}.xlsx"
    cert.xlsx_file.save(filename, ContentFile(patched), save=True)
    return cert.xlsx_file.url
